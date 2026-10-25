from dataclasses import dataclass, field
from pathlib import Path
import hashlib
import os
from typing import Any

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".tex",
    ".py", ".pyi", ".js", ".jsx", ".mjs", ".cjs",
    ".ts", ".tsx", ".java", ".kt", ".kts", ".c", ".h",
    ".cpp", ".hpp", ".cc", ".hh", ".cs", ".go", ".rs",
    ".php", ".rb", ".swift", ".scala", ".sh", ".bash",
    ".zsh", ".ps1", ".bat", ".cmd", ".sql", ".r", ".jl",
    ".lua", ".pl", ".html", ".htm", ".css", ".scss",
    ".sass", ".less", ".xml", ".svg", ".json", ".jsonl",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf",
    ".env", ".properties", ".graphql", ".gql", ".proto",
    ".log", ".csv", ".tsv",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
SPECIAL_TEXT_NAMES = {
    "dockerfile", "makefile", "justfile", ".gitignore",
    ".dockerignore", ".editorconfig", ".env",
    ".env.example", ".env.template",
}
IGNORED_DIRS = {
    ".git", ".venv", "venv", "__pycache__", "node_modules",
    "dist", "build", ".idea", ".vscode", ".pytest_cache",
    ".rag_cache",
}
MAX_FILE_BYTES = 50 * 1024 * 1024


@dataclass(frozen=True)
class SourceUnit:
    text: str
    start_char: int
    end_char: int
    locator: dict[str, Any] = field(default_factory=dict)
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class LoadedDocument:
    path: str
    content: str
    file_type: str
    sha256: str
    file_size: int
    units: list[SourceUnit]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError(f"Файл превышает лимит 50 МБ: {path}")
    if b"\x00" in raw:
        raise ValueError(f"Обнаружен бинарный файл: {path}")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    # latin-1 декодирует любые байты, поэтому не используем его как
    # безусловный fallback для потенциально бинарных файлов.
    raise UnicodeError(f"Не удалось декодировать текстовый файл: {path}")


def _units_from_blocks(blocks: list[tuple[str, dict, dict]]) -> tuple[str, list[SourceUnit]]:
    content_parts = []
    units = []
    offset = 0
    for text, locator, metadata in blocks:
        if not text or not text.strip():
            continue
        if content_parts:
            content_parts.append("\n\n")
            offset += 2
        start = offset
        content_parts.append(text)
        offset += len(text)
        units.append(SourceUnit(
            text=text,
            start_char=start,
            end_char=offset,
            locator=locator,
            metadata=metadata,
        ))
    return "".join(content_parts), units


def _load_pdf(path: Path) -> tuple[str, list[SourceUnit]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Для PDF установи зависимость: pip install pypdf") from exc

    reader = PdfReader(str(path))
    blocks = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append((
                text,
                {"kind": "pdf_page", "page": index},
                {"page": index},
            ))
    if not blocks:
        raise RuntimeError(
            f"В PDF не найден извлекаемый текст: {path.name}. "
            "Если это скан, понадобится OCR."
        )
    return _units_from_blocks(blocks)


def _load_docx(path: Path) -> tuple[str, list[SourceUnit]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Для DOCX установи зависимость: pip install python-docx") from exc

    document = Document(str(path))
    blocks = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            blocks.append((
                text,
                {"kind": "docx_paragraph", "paragraph_index": index},
                {"style": paragraph.style.name if paragraph.style else None},
            ))

    # Включаем таблицы: каждая строка таблицы становится отдельным блоком.
    table_index = 0
    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            text = " | ".join(value for value in cells if value)
            if text:
                blocks.append((
                    text,
                    {
                        "kind": "docx_table_row",
                        "table_index": table_index,
                        "row_index": row_index,
                    },
                    {},
                ))
        table_index += 1

    if not blocks:
        raise RuntimeError(f"В DOCX не найден текст: {path.name}")
    return _units_from_blocks(blocks)


def _load_image(path: Path) -> tuple[str, list[SourceUnit]]:
    try:
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError("Для изображений установи Pillow: pip install Pillow") from exc

    try:
        import pytesseract
    except ImportError as exc:
        raise RuntimeError(
            "Для OCR изображений установи pytesseract и Tesseract OCR. "
            "Команда Python: pip install pytesseract Pillow; "
            "сам Tesseract устанавливается отдельно."
        ) from exc

    with Image.open(path) as image:
        text = pytesseract.image_to_string(image, lang="rus+eng").strip()
        width, height = image.size

    if not text:
        raise RuntimeError(
            f"OCR не распознал текст в {path.name}. "
            "Для изображений без текста понадобится vision-модель."
        )

    blocks = [(
        text,
        {
            "kind": "image_ocr",
            "asset_path": str(path.resolve()),
            "bbox": [0, 0, width, height],
        },
        {"representation": "ocr", "width": width, "height": height},
    )]
    return _units_from_blocks(blocks)


def load_document(path: str | Path) -> LoadedDocument:
    path = Path(path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()
    name = path.name.lower()
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"Файл превышает лимит 50 МБ: {path}")

    if suffix in TEXT_EXTENSIONS or name in SPECIAL_TEXT_NAMES:
        content = _read_text_file(path)
        content, units = _units_from_blocks([
            (content, {
                "kind": "text_range",
                "start_char": 0,
                "end_char": len(content),
            }, {})
        ])
        file_type = suffix.lstrip(".") or name
    elif suffix == ".pdf":
        content, units = _load_pdf(path)
        file_type = "pdf"
    elif suffix == ".docx":
        content, units = _load_docx(path)
        file_type = "docx"
    elif suffix in IMAGE_EXTENSIONS:
        content, units = _load_image(path)
        file_type = suffix.lstrip(".")
    else:
        raise ValueError(f"Неподдерживаемый формат: {path.suffix or path.name}")

    return LoadedDocument(
        path=str(path),
        content=content,
        file_type=file_type,
        sha256=_sha256(path),
        file_size=path.stat().st_size,
        units=units,
    )


def load_documents(folder_path: str) -> list[LoadedDocument]:
    root = Path(folder_path).expanduser().resolve()
    if not root.is_dir():
        raise NotADirectoryError(f"Папка не существует: {root}")

    documents = []
    errors = []
    for current_dir, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(
            name for name in dirnames if name.lower() not in IGNORED_DIRS
        )
        for filename in sorted(filenames):
            path = Path(current_dir) / filename
            suffix = path.suffix.lower()
            if not (
                suffix in TEXT_EXTENSIONS
                or path.name.lower() in SPECIAL_TEXT_NAMES
                or suffix in {".pdf", ".docx"}
                or suffix in IMAGE_EXTENSIONS
            ):
                continue
            try:
                documents.append(load_document(path))
            except Exception as exc:
                errors.append(f"{path}: {exc}")

    if errors:
        preview = "\n".join(errors[:10])
        more = f"\n... и ещё {len(errors) - 10} ошибок" if len(errors) > 10 else ""
        raise RuntimeError(
            "Не удалось загрузить некоторые документы:\n"
            f"{preview}{more}\n\n"
            "Исправь ошибки или убери неподдерживаемые/повреждённые файлы "
            "из выбранной папки и повтори индексацию."
        )
    return documents
